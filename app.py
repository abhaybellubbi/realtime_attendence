from flask import Flask, render_template, request, jsonify, session, redirect, url_for, send_file
from functools import wraps
import os
import io
import csv
import threading
import time
from datetime import datetime, date
from werkzeug.security import check_password_hash
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

import database as db
import license_client
import excel_exporter
from biometric import BiometricConnector

app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'client_attendance_session_secret_771829')

# Initialize DB and Seed
db.init_db()
db.seed_default_data()

biometric_conn = BiometricConnector()

# ----------------- Decorators & Gatekeepers -----------------

def check_license_active():
    is_valid, status, info = license_client.check_license()
    return is_valid, status, info

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user' not in session:
            if request.path.startswith('/api/'):
                return jsonify({'error': 'Authentication required'}), 401
            return redirect(url_for('login_page'))
        return f(*args, **kwargs)
    return decorated_function

def require_super_admin(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user' not in session:
            return jsonify({'error': 'Authentication required'}), 401
        if session['user'].get('role') != 'super_admin':
            return jsonify({'error': 'Forbidden. Super Admin access required.'}), 403
        return f(*args, **kwargs)
    return decorated_function

def require_permission(perm_name):
    def decorator(f):
        @wraps(f)
        def decorated_function(*args, **kwargs):
            if 'user' not in session:
                return jsonify({'error': 'Authentication required'}), 401
            if session['user'].get('role') == 'worker':
                return jsonify({'error': 'Worker accounts cannot access administrative features'}), 403
            user_id = session['user']['id']
            if not db.has_permission(user_id, perm_name):
                return jsonify({
                    'error': f"Access denied. Missing permission: '{perm_name}'",
                    'permission_required': perm_name
                }), 403
            return f(*args, **kwargs)
        return decorated_function
    return decorator

# ----------------- App Lifecycle & License Verification -----------------

@app.before_request
def enforce_license_check():
    # Exempt static files and license activation routes
    exempt_paths = [
        '/static', '/login', '/logout', 
        '/api/license/reverify', '/api/license/activate', '/api/license/status'
    ]
    if any(request.path.startswith(p) for p in exempt_paths):
        return

    is_valid, status, info = check_license_active()
    if not is_valid:
        if request.path.startswith('/api/'):
            return jsonify({
                'error': 'Application license is inactive',
                'status': status,
                'message': info.get('message', 'License expired or suspended'),
                'company_name': info.get('company_name')
            }), 403
        # In web view, let index render with locked view
        return

@app.context_processor
def inject_template_globals():
    user = session.get('user', {})
    perms = {}
    if user and isinstance(user, dict):
        if user.get('role') == 'worker':
            perms = {'is_worker': True}
        elif user.get('id'):
            perms = db.get_user_permissions(user['id'])
    cached = license_client.get_cached_license() or {}
    lic_info = {
        'company_name': cached.get('company_name', 'Sai Krishna Enterprises'),
        'expiry_date': cached.get('expiry_date', '2027-01-01'),
        'days_remaining': cached.get('days_remaining', 90),
        'license_key': cached.get('license_key', 'LIC-APEX-8821-X9')
    }
    return {
        'user': user,
        'perms': perms,
        'license_info': lic_info,
        'license_valid': True,
        'license_status': 'ACTIVE'
    }

# ----------------- Auth & Main Routes -----------------

@app.route('/')
def home():
    if 'user' not in session:
        return redirect(url_for('login_page'))

    is_valid, status, lic_info = check_license_active()
    user = session['user']
    if user.get('role') == 'worker':
        perms = {'is_worker': True}
    else:
        perms = db.get_user_permissions(user['id'])

    return render_template(
        'index.html',
        user=user,
        perms=perms,
        departments=db.get_departments(),
        license_valid=is_valid,
        license_status=status,
        license_info=lic_info
    )

@app.route('/login', methods=['GET', 'POST'])
def login_page():
    if request.method == 'POST':
        data = request.get_json() if request.is_json else request.form
        username = data.get('username', '').strip()
        password = data.get('password', '')

        # 1. Check users table (super_admin, admin)
        user = db.get_user_by_username(username)
        if user and check_password_hash(user['password_hash'], password):
            perms = db.get_user_permissions(user['id'])
            u_dict = dict(user) if hasattr(user, 'keys') else user
            session['user'] = {
                'id': user['id'],
                'username': user['username'],
                'full_name': user['full_name'],
                'email': user['email'],
                'role': user['role'],
                'department_access': perms.get('department_access', 'ALL'),
                'phone': u_dict.get('phone') or ''
            }
            if request.is_json:
                return jsonify({'success': True, 'user': session['user'], 'permissions': perms})
            return redirect(url_for('home'))

        # 2. Check workers table (worker self-service)
        worker = db.authenticate_worker(username, password)
        if worker:
            w_dict = dict(worker) if hasattr(worker, 'keys') else worker
            session['user'] = {
                'id': worker['id'],
                'worker_id': worker['id'],
                'enroll_number': worker['enroll_number'],
                'username': f"worker_{worker['enroll_number']}",
                'full_name': worker['name'],
                'department': worker.get('department', 'General'),
                'shift_name': worker.get('shift_name', 'General Shift'),
                'email': worker.get('email', f"worker{worker['enroll_number']}@company.local"),
                'phone': w_dict.get('phone') or '',
                'role': 'worker'
            }
            if request.is_json:
                return jsonify({'success': True, 'user': session['user'], 'permissions': {'is_worker': True}})
            return redirect(url_for('home'))

        # 3. Fallback: Check if credentials belong to a client Super Admin created in Provider Portal
        lic_key_hint = data.get('license_key', '').strip() or None
        provider_auth = license_client.authenticate_with_provider(username, password, license_key=lic_key_hint)

        if provider_auth.get('authenticated'):
            # Auto-provision this Super Admin into local attendance database
            full_name = provider_auth.get('super_admin', {}).get('full_name') or f"{provider_auth['company_name']} (Owner)"
            email = f"{username}@{provider_auth.get('client_code', 'client').lower()}.com"
            db.provision_or_update_super_admin(username, password, full_name, email)

            # Auto-activate this client's license on the attendance app
            license_client.save_license_config(
                license_key=provider_auth['license_key'],
                company_name=provider_auth['company_name'],
                expiry_date=provider_auth['expiry_date'],
                status='active',
                days_remaining=provider_auth.get('days_remaining', 30)
            )

            user = db.get_user_by_username(username)
            perms = db.get_user_permissions(user['id'])
            u_dict = dict(user) if hasattr(user, 'keys') else user
            session['user'] = {
                'id': user['id'],
                'username': user['username'],
                'full_name': user['full_name'],
                'email': user['email'],
                'role': user['role'],
                'phone': u_dict.get('phone') or ''
            }
            if request.is_json:
                return jsonify({'success': True, 'user': session['user'], 'permissions': perms})
            return redirect(url_for('home'))
        else:
            err_msg = provider_auth.get('message', 'Invalid credentials')
            if request.is_json:
                return jsonify({'error': err_msg}), 401
            return render_template('index.html', show_login=True, login_error=err_msg)

    return render_template('index.html', show_login=True)

@app.route('/logout')
def logout():
    session.pop('user', None)
    return redirect(url_for('login_page'))

@app.route('/api/user', methods=['GET'])
@login_required
def api_current_user():
    user = session['user']
    if user.get('role') == 'worker':
        perms = {'is_worker': True}
    else:
        perms = db.get_user_permissions(user['id'])
    return jsonify({
        'user': user,
        'permissions': perms
    })

@app.route('/api/user/profile', methods=['POST'])
@login_required
def api_update_user_profile():
    user = session['user']
    data = request.get_json() or {}
    full_name = data.get('full_name', '').strip()
    email = data.get('email', '').strip()
    new_password = data.get('password', '').strip()
    current_password = data.get('current_password', '').strip()

    if not full_name:
        return jsonify({'error': 'Full name is required'}), 400

    if new_password:
        if not current_password:
            return jsonify({'error': 'Current password is required to change password'}), 400

        if user.get('role') == 'worker':
            worker_id = user.get('worker_id') or user.get('id')
            w = db.get_worker_by_enroll_number(user.get('enroll_number')) or db.get_worker_by_enroll_number(worker_id)
            if not w:
                return jsonify({'error': 'Worker record not found'}), 404
            stored_hash = w.get('password_hash')
            valid = False
            if stored_hash:
                valid = check_password_hash(stored_hash, current_password)
            else:
                valid = current_password in ('worker123', str(w.get('enroll_number')))
            if not valid:
                return jsonify({'error': 'Current password is incorrect'}), 400
        else:
            db_user = db.get_user_by_id(user['id'])
            if not db_user or not check_password_hash(db_user['password_hash'], current_password):
                return jsonify({'error': 'Current password is incorrect'}), 400

    phone = data.get('phone', '').strip()

    if user.get('role') == 'worker':
        worker_id = user.get('worker_id') or user.get('id')
        db.update_worker(worker_id, name=full_name, department=user.get('department', 'General'), email=email, phone=phone)
        if new_password:
            db.reset_worker_password(worker_id, new_password)
        session['user']['full_name'] = full_name
        session['user']['email'] = email
        session['user']['phone'] = phone
        session.modified = True
        return jsonify({'success': True, 'message': 'Profile updated successfully', 'user': session['user']})
    else:
        user_id = user['id']
        updated = db.update_user_profile(user_id, full_name, email=email, new_password=new_password if new_password else None, phone=phone)
        session['user']['full_name'] = updated['full_name']
        session['user']['email'] = updated['email']
        session['user']['phone'] = updated.get('phone') or phone
        session.modified = True
        return jsonify({'success': True, 'message': 'Profile updated successfully', 'user': session['user']})

# ----------------- License Management APIs -----------------

@app.route('/api/license/status', methods=['GET'])
def api_license_status():
    is_valid, status, info = license_client.check_license()
    cached = license_client.get_cached_license() or {}
    return jsonify({
        'valid': is_valid,
        'status': status,
        'company_name': cached.get('company_name', info.get('company_name', '')),
        'expiry_date': cached.get('expiry_date', info.get('expiry_date', '')),
        'days_remaining': cached.get('days_remaining', info.get('days_remaining', 0)),
        'license_key': cached.get('license_key', ''),
        'message': info.get('message', '')
    })

@app.route('/api/license/reverify', methods=['POST'])
def api_reverify_license():
    is_valid, status, info = license_client.check_license(force_online=True)
    return jsonify({
        'valid': is_valid,
        'status': status,
        'message': info.get('message', '')
    })

@app.route('/api/license/activate', methods=['POST'])
def api_activate_license():
    data = request.get_json() or {}
    key = data.get('license_key', '').strip()
    url = data.get('provider_url', 'http://localhost:5001/api/v1/license/verify').strip()

    if not key:
        return jsonify({'error': 'License key is required'}), 400

    license_client.save_license_config(license_key=key, provider_url=url)
    is_valid, status, info = license_client.check_license(force_online=True)
    if is_valid:
        return jsonify({'success': True, 'message': 'License activated successfully!', 'info': info})
    return jsonify({'error': info.get('message', 'License verification failed'), 'status': status}), 400

# ----------------- Customizable RBAC (Super Admin Only) -----------------

@app.route('/api/admins', methods=['GET'])
@require_super_admin
def api_get_admins():
    admins = db.get_all_admins()
    return jsonify({'admins': admins})

@app.route('/api/admins', methods=['POST'])
@require_super_admin
def api_create_admin():
    data = request.get_json() or {}
    username = data.get('username', '').strip()
    password = data.get('password', '').strip()
    full_name = data.get('full_name', '').strip()
    email = data.get('email', '').strip()
    permissions = data.get('permissions', {})

    if not username or not password or not full_name:
        return jsonify({'error': 'Username, password, and full name are required'}), 400

    success, res = db.create_admin_user(username, password, full_name, email, permissions)
    if success:
        return jsonify({'success': True, 'user_id': res, 'message': 'Admin account created successfully'}), 201
    return jsonify({'error': res}), 400

@app.route('/api/admins/<int:user_id>/permissions', methods=['PUT'])
@require_super_admin
def api_update_admin_permissions(user_id):
    data = request.get_json() or {}
    permissions = data.get('permissions', {})
    db.update_admin_permissions(user_id, permissions)
    return jsonify({'success': True, 'message': 'Permissions updated successfully'})

@app.route('/api/admins/<int:user_id>', methods=['DELETE'])
@require_super_admin
def api_delete_admin(user_id):
    db.delete_admin_user(user_id)
    return jsonify({'success': True, 'message': 'Admin account removed'})

# ----------------- Attendance & Reporting APIs (Protected by Granular RBAC) -----------------

@app.route('/api/stats', methods=['GET'])
@require_permission('can_view_dashboard')
def api_stats():
    return jsonify(db.get_stats())

@app.route('/api/departments', methods=['GET'])
@login_required
def api_get_departments():
    departments = db.get_departments()
    return jsonify({'success': True, 'departments': departments})

@app.route('/api/departments', methods=['POST'])
@login_required
def api_create_department():
    user = session.get('user', {})
    if user.get('role') not in ('super_admin', 'admin'):
        return jsonify({'error': 'Administrative privileges required to configure departments'}), 403

    data = request.get_json() or {}
    name = (data.get('name') or '').strip()
    code = (data.get('code') or '').strip()
    description = (data.get('description') or '').strip()

    if not name:
        return jsonify({'error': 'Department name is required'}), 400

    success, res = db.add_department(name, code=code, description=description)
    if success:
        return jsonify({
            'success': True,
            'message': f"Department '{name}' created successfully",
            'id': res
        }), 201
    return jsonify({'error': str(res)}), 400

@app.route('/api/departments/<int:dept_id>', methods=['PUT'])
@login_required
def api_update_department(dept_id):
    user = session.get('user', {})
    if user.get('role') not in ('super_admin', 'admin'):
        return jsonify({'error': 'Administrative privileges required to configure departments'}), 403

    data = request.get_json() or {}
    name = (data.get('name') or '').strip()
    code = (data.get('code') or '').strip()
    description = (data.get('description') or '').strip()

    if not name:
        return jsonify({'error': 'Department name cannot be empty'}), 400

    success, msg = db.update_department(dept_id, name, code=code, description=description)
    if success:
        return jsonify({'success': True, 'message': msg})
    return jsonify({'error': str(msg)}), 400

@app.route('/api/departments/<int:dept_id>', methods=['DELETE'])
@login_required
def api_delete_department(dept_id):
    user = session.get('user', {})
    if user.get('role') not in ('super_admin', 'admin'):
        return jsonify({'error': 'Administrative privileges required to delete departments'}), 403

    data = request.get_json(silent=True) or {}
    reassign_to = data.get('reassign_to') or request.args.get('reassign_to', 'General')

    success, msg = db.delete_department(dept_id, reassign_to=reassign_to)
    if success:
        return jsonify({'success': True, 'message': msg})
    return jsonify({'error': str(msg)}), 400

@app.route('/api/attendance/matrix', methods=['GET'])
@require_permission('can_view_attendance')
def api_attendance_matrix():
    year = int(request.args.get('year', 2026))
    month = int(request.args.get('month', 10))
    department = request.args.get('department')

    user = session.get('user', {})
    dept_access = user.get('department_access', 'ALL')
    if dept_access and dept_access != 'ALL':
        department = dept_access

    matrix = db.get_monthly_attendance_matrix(year, month, department)
    return jsonify(matrix)

@app.route('/api/attendance/daily', methods=['GET'])
@require_permission('can_view_attendance')
def api_attendance_daily():
    date_str = request.args.get('date', date.today().strftime('%Y-%m-%d'))
    daily = db.get_daily_attendance_details(date_str)
    return jsonify({'date': date_str, 'attendance': daily})

@app.route('/api/attendance/weekly', methods=['GET'])
@require_permission('can_view_attendance')
def api_attendance_weekly():
    date_str = request.args.get('date', date.today().strftime('%Y-%m-%d'))
    weekly = db.get_weekly_attendance_matrix(date_str)
    return jsonify(weekly)

@app.route('/api/attendance/export-weekly-excel', methods=['GET'])
@require_permission('can_export_excel')
def api_export_weekly_excel():
    date_str = request.args.get('date', date.today().strftime('%Y-%m-%d'))
    cached = license_client.get_cached_license() or {}
    company_name = cached.get('company_name', 'Sai Krishna Enterprises')

    weekly = db.get_weekly_attendance_matrix(date_str)
    excel_stream = excel_exporter.generate_weekly_excel(weekly, company_name)

    start_d = weekly.get('start_date', date_str)
    end_d = weekly.get('end_date', date_str)
    filename = f"{company_name.replace(' ', '_')}_Weekly_{start_d}_to_{end_d}.xlsx"
    return send_file(
        excel_stream,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

@app.route('/api/attendance/export-excel', methods=['GET'])
@require_permission('can_export_excel')
def api_export_excel():
    year = int(request.args.get('year', 2026))
    month = int(request.args.get('month', 10))
    department = request.args.get('department')
    mode = request.args.get('mode', 'single') # 'single' or 'department_pack'
    cached = license_client.get_cached_license() or {}
    company_name = cached.get('company_name', 'Sai Krishna Enterprises')

    user = session.get('user', {})
    dept_access = user.get('department_access', 'ALL')
    if dept_access and dept_access != 'ALL':
        department = dept_access
        mode = 'single'

    if mode == 'department_pack':
        full_matrix = db.get_monthly_attendance_matrix(year, month, department=None)
        excel_stream = excel_exporter.generate_department_pack_excel(full_matrix, company_name, year, month)
        filename = f"{company_name.replace(' ', '_')}_Monthly_Department_Pack_{year}_{month:02d}.xlsx"
    else:
        matrix = db.get_monthly_attendance_matrix(year, month, department)
        excel_stream = excel_exporter.generate_monthly_excel(matrix, company_name, year, month, department)
        dept_tag = f"_{department.replace(' ', '_')}" if department and department != 'ALL' else ""
        filename = f"{company_name.replace(' ', '_')}_Attendance{dept_tag}_{year}_{month:02d}.xlsx"

    return send_file(
        excel_stream,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=filename
    )

# ----------------- Workforce & Shift APIs -----------------

@app.route('/api/workers', methods=['GET'])
@require_permission('can_manage_workers')
def api_get_workers():
    workers = db.get_workers()
    return jsonify({'workers': workers})

@app.route('/api/workers', methods=['POST'])
@require_permission('can_manage_workers')
def api_create_worker():
    data = request.get_json() or {}
    enroll_number = int(data.get('enroll_number', 0))
    name = data.get('name', '').strip()
    department = data.get('department', 'General').strip()
    shift_id = data.get('shift_id')
    joining_date = data.get('joining_date')

    if not enroll_number or not name:
        return jsonify({'error': 'Enroll Number and Name are required'}), 400

    photo = data.get('photo', '').strip()
    success, res = db.create_worker(enroll_number, name, department, shift_id, joining_date, photo=photo)
    if success:
        return jsonify({'success': True, 'worker_id': res}), 201
    return jsonify({'error': res}), 400

@app.route('/api/workers/<int:worker_id>', methods=['PUT'])
@require_permission('can_manage_workers')
def api_update_worker(worker_id):
    data = request.get_json() or {}
    name = data.get('name', '').strip()
    department = data.get('department', 'General').strip()
    shift_id = data.get('shift_id')
    email = data.get('email')
    phone = data.get('phone')
    photo = data.get('photo')
    is_flexible = None
    if session['user'].get('role') == 'super_admin' and 'is_flexible' in data:
        is_flexible = bool(data.get('is_flexible'))
    db.update_worker(worker_id, name, department, shift_id, is_flexible=is_flexible, email=email, phone=phone, photo=photo)
    return jsonify({'success': True, 'message': 'Worker profile updated successfully'})

@app.route('/api/workers/<int:worker_id>/photo', methods=['POST'])
@require_permission('can_manage_workers')
def api_update_worker_photo(worker_id):
    data = request.get_json() or {}
    photo_data = data.get('photo', '').strip()
    db.update_worker_photo(worker_id, photo_data)
    return jsonify({'success': True, 'message': 'Worker photo updated successfully', 'photo': photo_data})

@app.route('/api/workers/<int:worker_id>/flexibility', methods=['POST'])
@login_required
def api_set_worker_flexibility(worker_id):
    if session['user'].get('role') != 'super_admin':
        return jsonify({'error': 'Permission denied: Only Super Admin can modify worker flexibility status'}), 403
    data = request.get_json() or {}
    is_flex = 1 if data.get('is_flexible') else 0
    db.set_worker_flexibility(worker_id, is_flex)
    return jsonify({'success': True, 'is_flexible': is_flex, 'message': f"Worker #{worker_id} set to {'Flexible Schedule' if is_flex else 'Fixed Shift'}."})

@app.route('/api/workers/<int:worker_id>/change-shift', methods=['POST'])
@require_permission('can_manage_workers')
def api_change_worker_shift(worker_id):
    data = request.get_json() or {}
    shift_id = data.get('shift_id')
    if shift_id is not None and str(shift_id).strip() != '':
        shift_id = int(shift_id)
    else:
        shift_id = None
    db.update_worker_shift(worker_id, shift_id)
    return jsonify({'success': True, 'message': 'Worker shift assignment updated successfully'})

@app.route('/api/workers/<int:worker_id>/leaves', methods=['GET'])
@require_permission('can_manage_leaves')
def api_get_worker_leaves_admin(worker_id):
    year = int(request.args.get('year', datetime.now().year))
    bal_data = db.get_worker_leave_balance(worker_id, year)
    balances = bal_data.get('balances', []) if isinstance(bal_data, dict) else bal_data
    return jsonify({'success': True, 'balances': balances, 'balance_summary': bal_data, 'year': year})

@app.route('/api/workers/<int:worker_id>/leaves/adjust', methods=['POST'])
@require_permission('can_manage_leaves')
def api_adjust_worker_leaves(worker_id):
    data = request.get_json() or {}
    leave_code = data.get('leave_code', '').strip().upper()
    year = int(data.get('year', datetime.now().year))
    allocated = data.get('allocated_days')
    used = data.get('used_days')
    add_allocated = data.get('add_allocated')
    add_used = data.get('add_used')

    if not leave_code:
        return jsonify({'error': 'Leave category code is required'}), 400

    ok, msg = db.adjust_worker_leave_balance(
        worker_id, leave_code, year,
        allocated_days=allocated,
        used_days=used,
        add_allocated=add_allocated,
        add_used=add_used
    )
    if ok:
        return jsonify({'success': True, 'message': msg})
    return jsonify({'error': msg}), 400

@app.route('/api/shifts/flexible-policy', methods=['GET'])
@login_required
def api_get_flexible_policy():
    policy = db.get_flexible_shift_policy()
    return jsonify({'success': True, 'policy': policy})

@app.route('/api/shifts/flexible-policy', methods=['POST'])
@require_permission('can_manage_shifts')
def api_save_flexible_policy():
    data = request.get_json() or {}
    try:
        full_day = float(data.get('full_day_hours', 8.0))
        half_day = float(data.get('half_day_hours', 4.0))
    except (ValueError, TypeError):
        return jsonify({'error': 'Full day and half day hours must be valid numbers'}), 400

    if full_day <= 0 or half_day <= 0:
        return jsonify({'error': 'Working hours thresholds must be greater than zero'}), 400
    if half_day > full_day:
        return jsonify({'error': 'Half day threshold cannot exceed full day threshold'}), 400

    updated = db.update_flexible_shift_policy(full_day, half_day)
    return jsonify({
        'success': True,
        'policy': updated,
        'message': f"Flexible shift policy updated: Full Day = {full_day} hrs, Half Day = {half_day} hrs."
    })

@app.route('/api/workers/<int:worker_id>', methods=['DELETE'])
@require_permission('can_manage_workers')
def api_delete_worker(worker_id):
    db.delete_worker(worker_id)
    return jsonify({'success': True, 'message': 'Worker deleted'})

@app.route('/api/shifts', methods=['GET'])
@login_required
def api_get_shifts():
    return jsonify({'shifts': db.get_shifts()})

@app.route('/api/shifts', methods=['POST'])
@require_permission('can_manage_shifts')
def api_create_shift():
    data = request.get_json() or {}
    name = data.get('name', '').strip()
    start_time = data.get('start_time', '').strip()
    end_time = data.get('end_time', '').strip()
    grace = int(data.get('grace_period_mins', 15))
    hd = int(data.get('half_day_threshold_mins', 120))

    if not name or not start_time or not end_time:
        return jsonify({'error': 'Name, Start Time, and End Time are required'}), 400

    shift_id = db.create_shift(name, start_time, end_time, grace, hd)
    return jsonify({'success': True, 'shift_id': shift_id}), 201

@app.route('/api/shifts/<int:shift_id>', methods=['PUT'])
@require_permission('can_manage_shifts')
def api_update_shift(shift_id):
    data = request.get_json() or {}
    name = data.get('name', '').strip()
    start_time = data.get('start_time', '').strip()
    end_time = data.get('end_time', '').strip()
    grace = int(data.get('grace_period_mins', 15))
    hd = int(data.get('half_day_threshold_mins', 120))

    db.update_shift(shift_id, name, start_time, end_time, grace, hd)
    return jsonify({'success': True, 'message': 'Shift updated'})

@app.route('/api/shifts/<int:shift_id>', methods=['DELETE'])
@require_permission('can_manage_shifts')
def api_delete_shift(shift_id):
    db.delete_shift(shift_id)
    return jsonify({'success': True, 'message': 'Shift deleted'})

# ----------------- Biometric Device Controls & SBXPC Integration -----------------

@app.route('/api/terminal/config', methods=['GET'])
@require_permission('can_trigger_sync')
def api_get_terminal_config():
    config = db.get_device_config(1)
    return jsonify({'success': True, 'config': config})

@app.route('/api/terminal/config', methods=['POST'])
@require_permission('can_trigger_sync')
def api_save_terminal_config():
    data = request.get_json() or {}
    ip = data.get('ip_address', '192.168.1.224').strip()
    port = int(data.get('port', 5005))
    mach_id = int(data.get('machine_id', 1))
    device_name = data.get('device_name', 'Front Counter Biometric Terminal').strip()
    sync_interval_mins = int(data.get('sync_interval_mins', 2))

    updated = db.update_device_config(
        1, ip_address=ip, port=port, machine_id=mach_id,
        device_name=device_name, sync_interval_mins=sync_interval_mins
    )
    return jsonify({
        'success': True,
        'message': f"Biometric terminal parameters saved (Auto-sync: every {sync_interval_mins} mins).",
        'config': updated
    })

@app.route('/api/terminal/test', methods=['POST'])
@require_permission('can_trigger_sync')
def api_test_terminal():
    data = request.get_json() or {}
    dev = db.get_device_config(1)
    ip = (data.get('ip_address') or (dev['ip_address'] if dev else '192.168.1.224')).strip()
    port = int(data.get('port') or (dev['port'] if dev else 5005))
    mach_id = int(data.get('machine_id') or (dev['machine_id'] if dev else 1))

    ok, msg = biometric_conn.connect(ip, port, mach_id)
    if ok:
        serial = biometric_conn.get_device_serial()
        if serial:
            db.update_device_config(1, serial_number=serial)
        biometric_conn.disconnect()
        db.log_sync(0, 'ONLINE', f"Hardware connection verified on {ip}:{port}")
        return jsonify({
            'success': True,
            'status': 'ONLINE',
            'message': f"Connected successfully to biometric hardware terminal at {ip}:{port}!",
            'serial_number': serial or (dev.get('serial_number') if dev else 'SBXPC-DEV-001')
        })
    else:
        db.log_sync(0, 'OFFLINE', f"Connection test failed: {msg}")
        return jsonify({
            'success': False,
            'status': 'OFFLINE',
            'message': f"Unable to reach device at {ip}:{port}. Error: {msg}"
        }), 503

@app.route('/api/terminal/sync-workers', methods=['POST'])
@require_permission('can_trigger_sync')
def api_sync_terminal_workers():
    dev = db.get_device_config(1)
    ip = dev['ip_address'] if dev else '192.168.1.224'
    port = dev['port'] if dev else 5005
    mach_id = dev['machine_id'] if dev else 1

    ok, msg = biometric_conn.connect(ip, port, mach_id)
    if not ok:
        db.log_sync(0, 'OFFLINE', f"Worker sync failed: {msg}")
        return jsonify({'success': False, 'status': 'OFFLINE', 'message': f"Cannot connect to hardware at {ip}:{port}: {msg}"}), 503

    ok_w, w_msg, w_list = biometric_conn.get_enrolled_workers()
    biometric_conn.disconnect()

    if not ok_w:
        db.log_sync(0, 'OFFLINE', f"Error reading enrolled users: {w_msg}")
        return jsonify({'success': False, 'message': w_msg}), 500

    sync_result = db.sync_enrolled_workers(w_list)
    res_msg = f"Biometric Sync: {sync_result['total']} worker records read from hardware. Added: {sync_result['added']} new, Updated: {sync_result['updated']} existing."
    db.log_sync(sync_result['added'], 'ONLINE', res_msg)

    return jsonify({
        'success': True,
        'message': res_msg,
        'workers_found': sync_result['total'],
        'new_workers_added': sync_result['added'],
        'workers_updated': sync_result['updated']
    })

@app.route('/api/terminal/sync', methods=['POST'])
@require_permission('can_trigger_sync')
def api_sync_terminal():
    dev = db.get_device_config(1)
    ip = dev['ip_address'] if dev else '192.168.1.224'
    port = dev['port'] if dev else 5005
    mach_id = dev['machine_id'] if dev else 1

    ok, msg = biometric_conn.connect(ip, port, mach_id)
    if not ok:
        db.log_sync(0, 'OFFLINE', f"Connection failed: {msg}")
        return jsonify({'success': False, 'status': 'OFFLINE', 'message': f"Cannot connect to hardware at {ip}:{port}: {msg}"}), 503

    # 1. Hardware RTC clock synchronization
    biometric_conn.sync_device_time()

    # 2. Sync enrolled workers and new employee IDs from hardware
    w_ok, w_msg, w_list = biometric_conn.get_enrolled_workers()
    worker_stats = {'added': 0, 'updated': 0, 'total': 0}
    if w_ok and w_list:
        worker_stats = db.sync_enrolled_workers(w_list)

    # 3. Read general attendance punch logs
    ok_logs, log_msg, logs = biometric_conn.get_attendance_logs(full_sync=True)
    biometric_conn.disconnect()

    if ok_logs:
        inserted = db.save_attendance_logs(1, logs)
        msg_summary = f"Full Biometric Sync Complete: {len(logs)} punch records fetched ({inserted} new punches). Enrolled workers: {worker_stats['total']} ({worker_stats['added']} newly enrolled)."
        db.log_sync(inserted, 'ONLINE', msg_summary)
        return jsonify({
            'success': True,
            'status': 'ONLINE',
            'synced': len(logs),
            'new_punches': inserted,
            'workers_synced': worker_stats['total'],
            'new_workers': worker_stats['added'],
            'message': msg_summary
        })
    else:
        db.log_sync(0, 'ONLINE', f"Workers synced ({worker_stats['total']}), punch read note: {log_msg}")
        return jsonify({
            'success': True,
            'status': 'ONLINE',
            'synced': 0,
            'new_punches': 0,
            'workers_synced': worker_stats['total'],
            'new_workers': worker_stats['added'],
            'message': f"Enrolled workers synced successfully ({worker_stats['total']}). Logs: {log_msg}"
        })

@app.route('/api/terminal/sync-clock', methods=['POST'])
@require_permission('can_trigger_sync')
def api_sync_terminal_clock():
    dev = db.get_device_config(1)
    ip = dev['ip_address'] if dev else '192.168.1.224'
    port = dev['port'] if dev else 5005
    mach_id = dev['machine_id'] if dev else 1

    ok, msg = biometric_conn.connect(ip, port, mach_id)
    if not ok:
        return jsonify({'success': False, 'message': f"Cannot connect to device at {ip}:{port}: {msg}"}), 503

    ok_clock, clock_msg = biometric_conn.sync_device_time()
    biometric_conn.disconnect()
    return jsonify({'success': ok_clock, 'message': clock_msg})

@app.route('/api/terminal/history', methods=['GET'])
@require_permission('can_trigger_sync')
def api_get_terminal_history():
    limit = int(request.args.get('limit', 15))
    history = db.get_sync_history(limit=limit)
    return jsonify({'success': True, 'history': history})

# ----------------- Worker Self-Service Endpoints -----------------

def worker_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user' not in session:
            return jsonify({'error': 'Authentication required'}), 401
        if session['user'].get('role') != 'worker':
            return jsonify({'error': 'Worker self-service portal is for worker accounts only'}), 403
        return f(*args, **kwargs)
    return decorated_function

@app.route('/api/worker/profile', methods=['GET'])
@worker_required
def api_worker_profile():
    worker_id = session['user'].get('worker_id') or session['user'].get('id')
    year = int(request.args.get('year', datetime.now().year))
    bal_data = db.get_worker_leave_balance(worker_id, year)
    balances = bal_data.get('balances', []) if isinstance(bal_data, dict) else bal_data
    return jsonify({
        'success': True,
        'worker': session['user'],
        'leave_balances': balances,
        'balance_summary': bal_data,
        'year': year
    })

@app.route('/api/worker/attendance', methods=['GET'])
@worker_required
def api_worker_attendance():
    worker_id = session['user'].get('worker_id') or session['user'].get('id')
    year = int(request.args.get('year', datetime.now().year))
    month = int(request.args.get('month', datetime.now().month))
    data = db.get_worker_personal_attendance(worker_id, year, month)
    return jsonify(data)

@app.route('/api/worker/leaves', methods=['GET'])
@worker_required
def api_worker_leaves():
    worker_id = session['user'].get('worker_id') or session['user'].get('id')
    year = int(request.args.get('year', datetime.now().year))
    bal_data = db.get_worker_leave_balance(worker_id, year)
    balances = bal_data.get('balances', []) if isinstance(bal_data, dict) else bal_data
    applications = db.get_leave_applications(worker_id=worker_id, year=year)
    leave_types = db.get_leave_types()
    return jsonify({
        'success': True,
        'balances': balances,
        'balance_summary': bal_data,
        'applications': applications,
        'leave_types': leave_types,
        'year': year
    })

@app.route('/api/worker/leaves', methods=['POST'])
@worker_required
def api_worker_apply_leave():
    worker_id = session['user'].get('worker_id') or session['user'].get('id')
    data = request.get_json() or {}
    leave_type_code = str(data.get('leave_type_code', '')).strip().upper()
    start_date = str(data.get('start_date', '')).strip()
    end_date = str(data.get('end_date', '')).strip()
    reason = str(data.get('reason', '')).strip()

    if not leave_type_code or not start_date or not end_date:
        return jsonify({'error': 'Leave type, start date, and end date are required'}), 400

    ok, res = db.apply_leave(
        worker_id=worker_id,
        leave_type_code=leave_type_code,
        start_date=start_date,
        end_date=end_date,
        reason=reason,
        applied_by='WORKER',
        applied_by_user_id=None
    )
    if ok:
        return jsonify({'success': True, 'leave_id': res, 'message': 'Leave application submitted successfully!'}), 201
    return jsonify({'error': res}), 400

# ----------------- Admin Leave Management Endpoints -----------------

@app.route('/api/admin/leaves', methods=['GET'])
@require_permission('can_manage_leaves')
def api_admin_get_leaves():
    status = request.args.get('status')
    if status == 'ALL' or not status:
        status = None
    year = int(request.args.get('year', datetime.now().year))
    month = request.args.get('month')
    month = int(month) if month and month.isdigit() else None

    applications = db.get_leave_applications(status=status, year=year, month=month)
    kpis = db.get_leave_kpis(year)
    leave_types = db.get_leave_types()

    return jsonify({
        'success': True,
        'applications': applications,
        'kpis': kpis,
        'leave_types': leave_types,
        'year': year
    })

@app.route('/api/admin/leaves/<int:leave_id>/review', methods=['POST'])
@require_permission('can_manage_leaves')
def api_admin_review_leave(leave_id):
    data = request.get_json() or {}
    action = str(data.get('status', '')).strip().upper()
    admin_notes = str(data.get('admin_notes', '')).strip()

    if action not in ['APPROVED', 'REJECTED']:
        return jsonify({'error': 'Status must be APPROVED or REJECTED'}), 400

    reviewer_id = session['user']['id']
    ok, msg = db.review_leave_application(leave_id, action, reviewer_id, admin_notes)
    if ok:
        return jsonify({'success': True, 'message': msg})
    return jsonify({'error': msg}), 400

@app.route('/api/admin/workers/<identifier>/leave-balance', methods=['GET'])
@require_permission('can_manage_leaves')
def api_admin_worker_leave_balance(identifier):
    worker = db.get_worker_by_enroll_or_id(identifier)
    if not worker:
        return jsonify({'error': f"Worker with Enroll ID or ID '{identifier}' not found"}), 404

    year = int(request.args.get('year', datetime.now().year))
    bal_data = db.get_worker_leave_balance(worker['id'], year)
    balances = bal_data.get('balances', []) if isinstance(bal_data, dict) else bal_data
    return jsonify({
        'success': True,
        'worker': worker,
        'balances': balances,
        'balance_summary': bal_data,
        'year': year
    })

@app.route('/api/admin/leaves/on-behalf', methods=['POST'])
@require_permission('can_manage_leaves')
def api_admin_leave_on_behalf():
    data = request.get_json() or {}
    identifier = str(data.get('identifier', '')).strip()
    leave_type_code = str(data.get('leave_type_code', '')).strip().upper()
    start_date = str(data.get('start_date', '')).strip()
    end_date = str(data.get('end_date', '')).strip()
    reason = str(data.get('reason', '')).strip()
    document_ref = str(data.get('document_ref', '')).strip()
    admin_notes = str(data.get('admin_notes', '')).strip()
    auto_approve = data.get('auto_approve', True)
    status = 'APPROVED' if auto_approve else 'PENDING'

    if not identifier or not leave_type_code or not start_date or not end_date:
        return jsonify({'error': 'Worker ID, Leave Type, Start Date, and End Date are required'}), 400

    worker = db.get_worker_by_enroll_or_id(identifier)
    if not worker:
        return jsonify({'error': f"Worker '{identifier}' not found in workforce"}), 404

    reviewer_id = session['user']['id'] if status == 'APPROVED' else None

    ok, res = db.apply_leave(
        worker_id=worker['id'],
        leave_type_code=leave_type_code,
        start_date=start_date,
        end_date=end_date,
        reason=reason,
        applied_by='ADMIN_ON_BEHALF',
        applied_by_user_id=session['user']['id'],
        document_ref=document_ref,
        status=status,
        reviewer_id=reviewer_id,
        admin_notes=admin_notes
    )
    if ok:
        return jsonify({
            'success': True,
            'leave_id': res,
            'message': f"Leave #{res} recorded & {status.lower()} for {worker['name']} (ID: {worker['enroll_number']})"
        }), 201
    return jsonify({'error': res}), 400

@app.route('/api/admin/leaves/balances', methods=['GET'])
@require_permission('can_manage_leaves')
def api_admin_leave_balances():
    year = int(request.args.get('year', datetime.now().year))
    balances = db.get_all_workers_leave_balances(year)
    return jsonify({'success': True, 'balances': balances, 'year': year})

# ----------------- Super Admin Leave Category Studio APIs -----------------

@app.route('/api/admin/leave-types', methods=['GET'])
@login_required
def api_admin_get_leave_types():
    include_inactive = request.args.get('include_inactive', 'false').lower() == 'true'
    types = db.get_leave_types(include_inactive=include_inactive)
    return jsonify({'success': True, 'leave_types': types})

@app.route('/api/admin/leave-types', methods=['POST'])
@require_super_admin
def api_admin_create_leave_type():
    data = request.get_json() or {}
    code = str(data.get('code', '')).strip().upper()
    name = str(data.get('name', '')).strip()
    default_days = data.get('default_days', 12)
    color = str(data.get('color', '#3b82f6')).strip()
    is_paid = 1 if data.get('is_paid', True) else 0
    description = str(data.get('description', '')).strip()

    if not code or not name:
        return jsonify({'error': 'Code and Name are required'}), 400

    ok, res = db.create_leave_type(code, name, default_days, color, is_paid, description)
    if ok:
        return jsonify({
            'success': True,
            'code': res,
            'message': f"Leave Category '{name}' ({res}) created and allocated to all active workers!"
        }), 201
    return jsonify({'error': res}), 400

@app.route('/api/admin/leave-types/<string:code>', methods=['PUT'])
@require_super_admin
def api_admin_update_leave_type(code):
    data = request.get_json() or {}
    name = str(data.get('name', '')).strip()
    default_days = data.get('default_days', 12)
    color = str(data.get('color', '#3b82f6')).strip()
    is_paid = 1 if data.get('is_paid', True) else 0
    description = str(data.get('description', '')).strip()
    update_workers = data.get('update_workers', True)
    year = int(data.get('year', datetime.now().year))

    if not name:
        return jsonify({'error': 'Category name cannot be empty'}), 400

    ok, res = db.update_leave_type(code, name, default_days, color, is_paid, description, update_workers, year)
    if ok:
        return jsonify({
            'success': True,
            'message': f"Leave category '{code}' updated successfully! New default: {default_days} days."
        })
    return jsonify({'error': res}), 400

@app.route('/api/admin/leave-types/<string:code>', methods=['DELETE'])
@require_super_admin
def api_admin_delete_leave_type(code):
    ok, res = db.delete_leave_type(code)
    if ok:
        return jsonify({'success': True, 'message': res})
    return jsonify({'error': res}), 400

# ----------------- Manual Check-In & Check-Out APIs -----------------

@app.route('/api/attendance/worker-status/<identifier>', methods=['GET'])
@login_required
def api_worker_punch_status(identifier):
    target_date = request.args.get('date', date.today().strftime('%Y-%m-%d'))
    status_data = db.get_worker_today_punch_status(identifier, target_date)
    if not status_data:
        return jsonify({'error': f"Worker '{identifier}' not found"}), 404
    return jsonify({'success': True, 'data': status_data})

@app.route('/api/attendance/manual-punch', methods=['POST'])
@login_required
def api_record_manual_punch():
    if session['user'].get('role') == 'worker':
        return jsonify({'error': 'Workers cannot perform manual biometric overrides'}), 403

    user_id = session['user']['id']
    if not (db.has_permission(user_id, 'can_manual_punch') or db.has_permission(user_id, 'can_edit_attendance')):
        return jsonify({'error': 'Permission denied: Manual Check-In/Out access required'}), 403

    data = request.get_json() or {}
    enroll_number = data.get('enroll_number')
    punch_type = int(data.get('punch_type', 0))
    p_date = data.get('date', date.today().strftime('%Y-%m-%d'))
    p_time = data.get('time', datetime.now().strftime('%H:%M:%S'))
    reason = data.get('reason', 'Manual Biometric Override')
    remarks = data.get('remarks', '')

    if not enroll_number:
        return jsonify({'error': 'Worker enroll number is required'}), 400

    # Ensure format YYYY-MM-DD HH:MM:SS
    if len(p_time) == 5:
        p_time += ":00"
    timestamp = f"{p_date} {p_time}"

    ok, res = db.record_manual_punch(
        enroll_number=int(enroll_number),
        punch_type=punch_type,
        event_timestamp=timestamp,
        recorded_by_user_id=user_id,
        reason=reason,
        remarks=remarks
    )
    if ok:
        return jsonify({'success': True, 'data': res}), 201
    return jsonify({'error': res}), 400

@app.route('/api/attendance/manual-punches', methods=['GET'])
@login_required
def api_get_manual_punches():
    target_date = request.args.get('date')
    limit = int(request.args.get('limit', 50))
    punches = db.get_manual_punches(limit=limit, target_date=target_date)
    return jsonify({'success': True, 'punches': punches})

@app.route('/api/attendance/manual-punch/<int:punch_id>', methods=['DELETE'])
@login_required
def api_delete_manual_punch(punch_id):
    if session['user'].get('role') == 'worker':
        return jsonify({'error': 'Workers cannot delete punches'}), 403

    user_id = session['user']['id']
    if not (db.has_permission(user_id, 'can_manual_punch') or db.has_permission(user_id, 'can_edit_attendance')):
        return jsonify({'error': 'Permission denied'}), 403

    ok, res = db.delete_manual_punch(punch_id)
    if ok:
        return jsonify({'success': True, 'message': res})
    return jsonify({'error': res}), 400


# ----------------- Background Worker -----------------

def background_sync_worker():
    # Initial pause to allow Flask server to bind cleanly
    time.sleep(3)
    last_sync_timestamp = 0

    while True:
        try:
            dev = db.get_device_config(1)
            interval_mins = dev.get('sync_interval_mins', 2) if dev else 2
            interval_secs = max(60, int(interval_mins * 60))

            now_ts = time.time()
            if (now_ts - last_sync_timestamp) >= interval_secs:
                last_sync_timestamp = now_ts
                if dev and dev.get('is_active', 1):
                    ok, msg = biometric_conn.connect(dev['ip_address'], dev['port'], dev['machine_id'])
                    if ok:
                        # Sync workers enrolled directly on hardware
                        w_ok, w_msg, w_list = biometric_conn.get_enrolled_workers()
                        if w_ok and w_list:
                            db.sync_enrolled_workers(w_list)

                        # Sync attendance punches
                        ok_logs, log_msg, logs = biometric_conn.get_attendance_logs(full_sync=False)
                        biometric_conn.disconnect()

                        if ok_logs and logs:
                            inserted = db.save_attendance_logs(1, logs)
                            db.log_sync(inserted, 'ONLINE', f"Auto background sync ({interval_mins}m): {len(logs)} punches ({inserted} new)")
                        else:
                            db.log_sync(0, 'ONLINE', f"Auto background sync ({interval_mins}m) completed (no new punches)")
                    else:
                        db.log_sync(0, 'OFFLINE', f"Auto sync: {msg}")

            time.sleep(5)
        except Exception as e:
            print(f"[Worker] Background sync exception: {e}")
            time.sleep(10)

sync_thread = threading.Thread(target=background_sync_worker, daemon=True)
sync_thread.start()


# ----------------- Official Holidays Management APIs -----------------

@app.route('/api/holidays', methods=['GET'])
@login_required
def api_get_holidays():
    year = request.args.get('year')
    year_int = int(year) if year and year.isdigit() else None
    holidays = db.get_holidays(year=year_int)
    return jsonify({'success': True, 'holidays': holidays})

@app.route('/api/holidays', methods=['POST'])
@login_required
def api_add_holiday():
    if session['user'].get('role') != 'super_admin':
        return jsonify({'error': 'Permission denied: Only Super Admin can configure official holidays'}), 403

    data = request.get_json() or {}
    holiday_date = (data.get('holiday_date') or data.get('date') or '').strip()
    holiday_name = (data.get('holiday_name') or data.get('name') or '').strip()
    notes = (data.get('notes') or '').strip()

    if not holiday_date or not holiday_name:
        return jsonify({'error': 'Holiday date and holiday name are required'}), 400

    user_id = session['user']['id']
    ok, res = db.add_holiday(holiday_date, holiday_name, notes=notes, user_id=user_id)
    if ok:
        return jsonify({'success': True, 'holiday_id': res, 'message': f"Holiday '{holiday_name}' registered successfully."}), 201
    return jsonify({'error': res}), 400

@app.route('/api/holidays/<int:holiday_id>', methods=['PUT'])
@login_required
def api_update_holiday(holiday_id):
    if session['user'].get('role') != 'super_admin':
        return jsonify({'error': 'Permission denied: Only Super Admin can modify official holidays'}), 403

    data = request.get_json() or {}
    holiday_date = data.get('holiday_date', '').strip()
    holiday_name = data.get('holiday_name', '').strip()
    notes = data.get('notes', '').strip()

    ok, res = db.update_holiday(holiday_id, holiday_date, holiday_name, notes=notes)
    if ok:
        return jsonify({'success': True, 'message': 'Holiday updated successfully.'})
    return jsonify({'error': res}), 400

@app.route('/api/holidays/<int:holiday_id>', methods=['DELETE'])
@login_required
def api_delete_holiday(holiday_id):
    if session['user'].get('role') != 'super_admin':
        return jsonify({'error': 'Permission denied: Only Super Admin can delete official holidays'}), 403

    deleted = db.delete_holiday(holiday_id)
    if deleted:
        return jsonify({'success': True, 'message': 'Holiday removed successfully.'})
    return jsonify({'error': 'Holiday not found or could not be removed'}), 404

@app.route('/api/holidays/export', methods=['GET'])
@login_required
def api_export_holidays():
    year = request.args.get('year')
    year_int = int(year) if year and year.isdigit() else None
    holidays = db.get_holidays(year=year_int)

    wb = openpyxl.Workbook()
    ws = wb.active
    year_label = str(year_int) if year_int else "All_Years"
    ws.title = f"Holidays_{year_label}"
    ws.views.sheetView[0].showGridLines = True

    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title Banner
    ws.merge_cells('A1:E1')
    ws['A1'] = "SAI KRISHNA ENTERPRISES - OFFICIAL GOVERNMENT & COMPANY HOLIDAYS"
    ws['A1'].font = Font(name="Calibri", size=15, bold=True, color="0F172A")
    ws['A1'].alignment = Alignment(horizontal="center", vertical="center")

    ws.merge_cells('A2:E2')
    ws['A2'] = f"Official Calendar Schedule ({year_label}) | Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    ws['A2'].font = Font(name="Calibri", size=10, italic=True, color="64748B")
    ws['A2'].alignment = Alignment(horizontal="center", vertical="center")

    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[4].height = 26

    cols = ["Holiday Date (YYYY-MM-DD)", "Day of Week", "Holiday Name", "Guidelines & Notes", "Attendance Entitlement"]
    for col_idx, col_name in enumerate(cols, 1):
        cell = ws.cell(row=4, column=col_idx, value=col_name)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    current_row = 5
    for h in holidays:
        ws.row_dimensions[current_row].height = 22

        c1 = ws.cell(row=current_row, column=1, value=h.get('holiday_date'))
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.font = Font(name="Calibri", size=10, bold=True, color="4338CA")
        c1.border = thin_border

        c2 = ws.cell(row=current_row, column=2, value=h.get('day_of_week'))
        c2.alignment = Alignment(horizontal="center", vertical="center")
        c2.font = Font(name="Calibri", size=10)
        c2.border = thin_border

        c3 = ws.cell(row=current_row, column=3, value=h.get('holiday_name'))
        c3.alignment = Alignment(horizontal="left", vertical="center")
        c3.font = Font(name="Calibri", size=10, bold=True)
        c3.border = thin_border

        c4 = ws.cell(row=current_row, column=4, value=h.get('notes') or 'Official Public Holiday')
        c4.alignment = Alignment(horizontal="left", vertical="center")
        c4.font = Font(name="Calibri", size=9, color="475569")
        c4.border = thin_border

        c5 = ws.cell(row=current_row, column=5, value="Full Paid Holiday (No Absent Penalty)")
        c5.alignment = Alignment(horizontal="center", vertical="center")
        c5.font = Font(name="Calibri", size=9, bold=True, color="059669")
        c5.border = thin_border

        current_row += 1

    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 38
    ws.column_dimensions['D'].width = 44
    ws.column_dimensions['E'].width = 35

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(
        out,
        as_attachment=True,
        download_name=f"Official_Holidays_{year_label}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route('/api/holidays/template', methods=['GET'])
@login_required
def api_holidays_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Holidays_Template"
    ws.views.sheetView[0].showGridLines = True

    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    cols = ["Holiday Date (YYYY-MM-DD)", "Holiday Name", "Guidelines / Notes"]
    for col_idx, col_name in enumerate(cols, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    sample_rows = [
        ("2026-01-26", "Republic Day (ಗಣರಾಜ್ಯೋತ್ಸವ)", "National Gazetted Holiday"),
        ("2026-03-20", "Ugadi (ಯುಗಾದಿ ಹಬ್ಬ)", "Karnataka State New Year Festival"),
        ("2026-05-01", "May Day (ಕಾರ್ಮಿಕರ ದಿನಾಚರಣೆ)", "International Workers Day - Mandatory Paid Holiday"),
        ("2026-08-15", "Independence Day (ಸ್ವಾತಂತ್ರ್ಯ ದಿನಾಚರಣೆ)", "National Gazetted Holiday"),
        ("2026-10-02", "Mahatma Gandhi Jayanti (ಗಾಂಧಿ ಜಯಂತಿ)", "Father of the Nation Birthday"),
        ("2026-11-01", "Kannada Rajyotsava (ಕನ್ನಡ ರಾಜ್ಯೋತ್ಸವ)", "Karnataka State Formation Day"),
        ("2026-12-25", "Christmas Day (ಕ್ರಿಸ್‌ಮಸ್)", "Public Gazetted Holiday")
    ]

    for row_idx, r_data in enumerate(sample_rows, 2):
        for col_idx, val in enumerate(r_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = thin_border
            cell.font = Font(name="Calibri", size=10)

    ws.column_dimensions['A'].width = 28
    ws.column_dimensions['B'].width = 38
    ws.column_dimensions['C'].width = 44

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(
        out,
        as_attachment=True,
        download_name="Holidays_Import_Template.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

@app.route('/api/holidays/import', methods=['POST'])
@login_required
def api_import_holidays():
    if session['user'].get('role') != 'super_admin':
        return jsonify({'error': 'Permission denied: Only Super Admin can import official holidays'}), 403

    holidays_to_import = []

    # Check for direct JSON payload
    if request.is_json:
        data = request.get_json() or {}
        holidays_to_import = data.get('holidays', [])
    elif 'file' in request.files:
        uploaded_file = request.files['file']
        if not uploaded_file or not uploaded_file.filename:
            return jsonify({'error': 'No file was selected for upload.'}), 400

        fname = uploaded_file.filename.lower()
        if fname.endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
            try:
                wb = openpyxl.load_workbook(uploaded_file, data_only=True)
                ws = wb.active
                for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                    if row_idx == 1:
                        # Skip header row if it contains 'Date' or 'Holiday'
                        first_val = str(row[0] or '').lower()
                        if 'date' in first_val or 'holiday' in first_val or 'day' in first_val:
                            continue

                    if not row or not row[0]:
                        continue

                    # Handle date field (string or datetime)
                    raw_date = row[0]
                    if isinstance(raw_date, (datetime, date)):
                        date_str = raw_date.strftime('%Y-%m-%d')
                    else:
                        date_str = str(raw_date).strip()[:10]

                    name_str = str(row[1] or '').strip()
                    notes_str = str(row[2] or '').strip() if len(row) > 2 else ''

                    if date_str and name_str:
                        holidays_to_import.append({
                            'holiday_date': date_str,
                            'holiday_name': name_str,
                            'notes': notes_str
                        })
            except Exception as e:
                return jsonify({'error': f"Failed to parse Excel file: {str(e)}"}), 400
        elif fname.endswith('.csv'):
            try:
                stream = io.StringIO(uploaded_file.stream.read().decode('utf-8-sig', errors='replace'))
                reader = csv.reader(stream)
                for row_idx, row in enumerate(reader, 1):
                    if row_idx == 1:
                        first_val = str(row[0] or '').lower() if row else ''
                        if 'date' in first_val or 'holiday' in first_val:
                            continue
                    if len(row) >= 2 and row[0].strip() and row[1].strip():
                        holidays_to_import.append({
                            'holiday_date': row[0].strip()[:10],
                            'holiday_name': row[1].strip(),
                            'notes': row[2].strip() if len(row) > 2 else ''
                        })
            except Exception as e:
                return jsonify({'error': f"Failed to parse CSV file: {str(e)}"}), 400
        else:
            return jsonify({'error': 'Unsupported file format. Please upload an Excel (.xlsx) or CSV (.csv) file.'}), 400

    if not holidays_to_import:
        return jsonify({'error': 'No valid holiday records found in the file. Ensure Date and Holiday Name are provided.'}), 400

    user_id = session['user']['id']
    ok, ins, upd = db.bulk_import_holidays(holidays_to_import, user_id=user_id)
    if ok:
        return jsonify({
            'success': True,
            'message': f"Import complete! {ins} new holidays added, {upd} existing holidays updated.",
            'inserted': ins,
            'updated': upd,
            'total': ins + upd
        })
    return jsonify({'error': str(ins)}), 400

@app.route('/api/holidays/seed-karnataka', methods=['POST'])
@login_required
def api_seed_karnataka_holidays():
    if session['user'].get('role') != 'super_admin':
        return jsonify({'error': 'Permission denied: Only Super Admin can load national/state holidays'}), 403

    user_id = session['user']['id']
    ok, ins, upd = db.seed_karnataka_national_holidays(user_id=user_id)
    if ok:
        return jsonify({
            'success': True,
            'message': f"Loaded official Karnataka & National holidays! ({ins} new added, {upd} existing updated)",
            'inserted': ins,
            'updated': upd
        })
    return jsonify({'error': str(ins)}), 400

@app.route('/api/attendance/weekoff-compoffs', methods=['GET'])
@login_required
def api_weekoff_compoffs():
    year = request.args.get('year')
    year_int = int(year) if year and year.isdigit() else datetime.now().year
    summary = db.get_weekoff_workers_summary(year=year_int)
    return jsonify({'success': True, 'year': year_int, 'data': summary})

@app.route('/api/attendance/sync-compoffs', methods=['POST'])
@login_required
def api_sync_compoffs():
    if session['user'].get('role') not in ('super_admin', 'admin'):
        return jsonify({'error': 'Permission denied: Administrative role required to sync comp-offs'}), 403

    year = request.args.get('year')
    year_int = int(year) if year and year.isdigit() else datetime.now().year
    updated = db.sync_weekoff_earned_compoffs(year=year_int)
    return jsonify({
        'success': True,
        'message': f"Successfully synchronized weekoff comp-offs for {year_int}. Updated {updated} worker balance records.",
        'updated_count': updated
    })

if __name__ == '__main__':
    print("=" * 60)
    print("  SAI KRISHNA ENTERPRISES - WORKFORCE ATTENDANCE SYSTEM")
    print("  URL: http://localhost:5000")
    print("  Status: Active & Secure")
    print("=" * 60)
    app.run(host='0.0.0.0', port=5000, debug=True)
